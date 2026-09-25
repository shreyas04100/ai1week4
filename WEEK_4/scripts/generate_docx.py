"""
generate_docx.py — Week 4 Task 4
Generates: Task_4_Performance_Evaluation_Impact_Framework_NovaMart_FINAL.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT = os.path.join(os.path.dirname(__file__), "..", "documents",
                      "Task_4_Performance_Evaluation_Impact_Framework_NovaMart_FINAL.docx")

# ── Palette ──────────────────────────────────────────────────────────────────
NAVY   = RGBColor(0x1F, 0x35, 0x64)
TEAL   = RGBColor(0x1F, 0x7A, 0x8C)
ORANGE = RGBColor(0xE8, 0x6A, 0x1A)
GREEN  = RGBColor(0x1A, 0x7A, 0x3C)
PURPLE = RGBColor(0x6A, 0x1A, 0x7A)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY  = RGBColor(0xF2, 0xF2, 0xF2)
MGRAY  = RGBColor(0xC0, 0xC0, 0xC0)
DGRAY  = RGBColor(0x40, 0x40, 0x40)
LBLUE  = RGBColor(0xD6, 0xE4, 0xF0)
LGREEN = RGBColor(0xD5, 0xF0, 0xE0)
LYELL  = RGBColor(0xFD, 0xF6, 0xD3)
LRED   = RGBColor(0xF9, 0xE0, 0xE0)
LPURP  = RGBColor(0xEE, 0xE0, 0xF8)


def _hex(rgb):
    return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"


def set_cell_bg(cell, rgb):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex(rgb))
    tcPr.append(shd)


def add_run(para, text, bold=False, italic=False,
            size=11, color=DGRAY, font="Calibri"):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return run


def heading(doc, text, level=1):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    para.paragraph_format.space_after = Pt(4)
    if level == 1:
        add_run(para, text, bold=True, size=15, color=NAVY)
    elif level == 2:
        add_run(para, text, bold=True, size=12, color=TEAL)
    else:
        add_run(para, text, bold=True, size=11, color=ORANGE)
    return para


def body(doc, text, size=10.5, italic=False, space_after=5):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    add_run(para, text, size=size, italic=italic)
    return para


def bullet(doc, text, size=10.5):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(2)
    add_run(para, text, size=size)
    return para


def kv(doc, key, val, size=10.5):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    add_run(para, key + ": ", bold=True, size=size, color=TEAL)
    add_run(para, val, size=size)


def formula(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.left_indent = Inches(0.4)
    add_run(para, text, bold=True, size=10.5, color=PURPLE, font="Courier New")


def note(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.left_indent = Inches(0.3)
    add_run(para, "⚠ Note: ", bold=True, size=9.5, color=ORANGE)
    add_run(para, text, size=9.5, italic=True)


def divider(doc):
    para = doc.add_paragraph()
    pPr = para._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "6")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), "1F7A8C")
    pb.append(bot)
    pPr.append(pb)


def simple_table(doc, headers, rows, col_widths=None, hdr_color=None):
    hdr_color = hdr_color or NAVY
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hrow = table.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        set_cell_bg(cell, hdr_color)
        p = cell.paragraphs[0]
        add_run(p, h, bold=True, size=9, color=WHITE)
    for ri, row in enumerate(rows):
        bg = LGRAY if ri % 2 == 0 else WHITE
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            add_run(p, str(val), size=9)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def flow_diagram(doc, boxes, title, cols=3):
    """Build a vertical flow diagram from a list of (label, color) tuples."""
    n = len(boxes)
    rows_count = n * 2 - 1
    tbl = doc.add_table(rows=rows_count, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in range(rows_count):
        for c in range(cols):
            cell = tbl.cell(r, c)
            set_cell_bg(cell, WHITE)
            cell.paragraphs[0].text = ""
            cell.width = Inches(2.0)
    for i, (label, color) in enumerate(boxes):
        row_idx = i * 2
        cell = tbl.cell(row_idx, 1)
        set_cell_bg(cell, color)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        add_run(p, label, bold=True, size=9, color=WHITE)
        if i < n - 1:
            arr = tbl.cell(row_idx + 1, 1)
            set_cell_bg(arr, WHITE)
            ap = arr.paragraphs[0]
            ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run(ap, "↓", bold=True, size=12, color=TEAL)
    doc.add_paragraph()
    body(doc, title, italic=True, size=9.5)
    return tbl


def add_header_footer(doc):
    for section in doc.sections:
        header = section.header
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp.clear()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_run(hp, "NovaMart | Performance Evaluation & Impact Assessment | Week 4 Task 4",
                size=8, italic=True, color=MGRAY)
        footer = section.footer
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.clear()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(fp, "CONFIDENTIAL — Hypothetical Company — Internship Exercise Only",
                size=8, italic=True, color=MGRAY)


# ── SECTION BUILDERS ─────────────────────────────────────────────────────────

def cover_page(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "NOVAMART", bold=True, size=30, color=NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p2, "Performance Evaluation and Impact Assessment Framework",
            bold=True, size=18, color=TEAL)
    doc.add_paragraph()
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p3, "AI Initiative Evaluation Framework for NovaMart Retail & E-Commerce Pvt. Ltd.",
            bold=True, size=13, color=ORANGE)
    doc.add_paragraph()
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p4, "Internship Week 4 — Task 4 Deliverable", size=12, color=DGRAY)
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p5, "September 2026", size=11, italic=True, color=DGRAY)
    doc.add_paragraph()
    p6 = doc.add_paragraph()
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p6,
            "DISCLAIMER: NovaMart Retail & E-Commerce Private Limited is a hypothetical company "
            "created for this internship exercise. All KPI targets, financial projections, case "
            "study results, and evaluation examples are illustrative planning assumptions and "
            "hypothetical scenarios. They do not represent actual NovaMart data or results.",
            size=9, italic=True, color=RGBColor(0x80, 0x80, 0x80))
    doc.add_page_break()


def exec_summary(doc):
    heading(doc, "1. Executive Summary", 1)
    divider(doc)
    body(doc,
         "This document presents a comprehensive Performance Evaluation and Impact Assessment "
         "Framework for NovaMart Retail & E-Commerce Private Limited (hypothetical). The "
         "framework provides a structured, multi-dimensional methodology for measuring the "
         "performance, business impact, financial value, and governance of AI initiatives "
         "deployed across NovaMart's retail operations.")

    heading(doc, "Why AI Performance Must Be Measured", 2)
    body(doc,
         "Deploying an AI system is not the end of the journey — it is the beginning. "
         "Without rigorous measurement, NovaMart cannot determine whether AI investments "
         "are delivering value, identify underperforming models before they cause business "
         "harm, justify continued AI investment to leadership, or continuously improve "
         "AI systems based on real-world feedback.")

    heading(doc, "The Critical Distinction: Model Performance vs Business Impact", 2)
    body(doc,
         "A fundamental principle of this framework is that technical model performance "
         "and business impact are NOT the same thing and must be measured separately:")
    simple_table(doc,
        ["Dimension", "What It Measures", "Example", "Who Cares"],
        [
            ["AI Model Performance", "How accurately the model predicts", "MAPE = 11% for demand forecasting", "Data Science team"],
            ["System Performance", "How reliably the system operates", "API uptime 99.7%, latency <200ms", "Engineering team"],
            ["Operational Performance", "How efficiently business processes run", "Replenishment cycle time reduced 60%", "Operations team"],
            ["Business Impact", "How much business value was created", "Stockout rate reduced from 15% to 8%", "Business unit heads"],
            ["Financial Impact", "How much money was made or saved", "INR 12 Cr annual inventory cost reduction", "CFO / Finance"],
            ["Customer Impact", "How customers were affected", "NPS improved from 42 to 58", "CX / Marketing"],
            ["Strategic Impact", "How competitive position changed", "First-mover AI advantage in Tier-2 markets", "CEO / Board"],
        ],
        col_widths=[1.5, 1.8, 2.0, 1.2]
    )
    note(doc, "A model can have excellent technical accuracy but still produce poor business "
         "results if users do not adopt its recommendations, if data quality is poor, or if "
         "the business process around the AI is not redesigned to act on AI outputs.")

    heading(doc, "Framework Summary", 2)
    bullet(doc, "9 evaluation dimensions: Model, Data, System, Operational, Customer, Financial, Adoption, Risk/Governance, Strategic.")
    bullet(doc, "5 detailed case study scenarios: Demand Forecasting, Personalization, Fraud Detection, Churn Prediction, GenAI Assistant.")
    bullet(doc, "Structured ROI methodology with direct, indirect, tangible, and intangible benefit categories.")
    bullet(doc, "Proposed AI Performance Dashboard with 8 panels.")
    bullet(doc, "30/60/90-day evaluation plan and 18-month long-term roadmap.")
    doc.add_page_break()


def introduction(doc):
    heading(doc, "2. Introduction", 1)
    divider(doc)
    body(doc,
         "NovaMart (hypothetical) has embarked on a multi-year AI transformation program "
         "spanning demand forecasting, customer personalization, fraud detection, churn "
         "prediction, and generative AI-powered customer service. As these AI initiatives "
         "move from development into production, a rigorous evaluation framework becomes "
         "essential to ensure that AI investments deliver measurable, sustainable business value.")

    heading(doc, "3. NovaMart Business Context", 1)
    divider(doc)
    body(doc,
         "NovaMart operates 200+ physical stores, an e-commerce platform, and 20 dark stores "
         "across Indian Tier-1, Tier-2, and Tier-3 cities (illustrative). The company manages "
         "~40,000–50,000 active SKUs and serves 2M+ registered customers. AI initiatives "
         "evaluated in this framework are drawn from the Week 2 opportunity analysis and "
         "the Week 3 ADFIO architecture.")

    simple_table(doc,
        ["AI Initiative", "Week Reference", "Primary Business Goal"],
        [
            ["AI Demand Forecasting & Inventory Optimization (ADFIO)", "Week 3 Architecture", "Reduce stockouts and inventory costs"],
            ["Real-Time Customer Personalization", "Week 2 Opportunity #1", "Increase conversion and basket size"],
            ["AI Fraud Detection", "Week 2 Opportunity #6", "Reduce fraud losses and false positives"],
            ["Customer Churn Prediction", "Week 2 Opportunity #5", "Retain high-value customers"],
            ["GenAI Conversational Commerce Assistant", "Week 2 Opportunity #15", "Deflect support tickets; improve CX"],
        ],
        col_widths=[2.8, 1.5, 2.2]
    )

    heading(doc, "4. Purpose of the Evaluation Framework", 1)
    divider(doc)
    bullet(doc, "Provide a structured, repeatable methodology for evaluating all NovaMart AI initiatives.")
    bullet(doc, "Connect technical AI metrics to business KPIs and financial outcomes.")
    bullet(doc, "Enable data-driven decisions about AI investment, continuation, scaling, or discontinuation.")
    bullet(doc, "Support continuous improvement of AI models and business processes.")
    bullet(doc, "Ensure AI governance, fairness, and compliance with DPDP Act 2023.")
    bullet(doc, "Provide transparent reporting to business stakeholders and leadership.")

    heading(doc, "5. Evaluation Principles", 1)
    divider(doc)
    simple_table(doc,
        ["Principle", "Description"],
        [
            ["Baseline First", "No AI impact can be measured without a pre-deployment baseline. Baselines must be established and signed off before AI goes live."],
            ["Multi-Dimensional", "Evaluate across technical, operational, financial, customer, and governance dimensions simultaneously."],
            ["Causality over Correlation", "Use controlled experiments (A/B tests) or rigorous before/after analysis to establish causality, not just correlation."],
            ["Proportionality", "Measurement effort should be proportional to the business value and risk of the AI initiative."],
            ["Transparency", "All assumptions, limitations, and uncertainties in measurement must be clearly documented."],
            ["Actionability", "Every metric must have a defined owner, threshold, and action to take when performance is poor."],
            ["Continuous", "Evaluation is not a one-time activity. It is an ongoing process embedded in the AI lifecycle."],
            ["Fairness", "Evaluation must include bias and fairness checks to ensure AI does not discriminate against customer segments."],
        ],
        col_widths=[1.5, 5.0]
    )
    doc.add_page_break()
